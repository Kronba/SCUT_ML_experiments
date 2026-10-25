# -*- coding: utf-8 -*-
"""docx 排版工具箱：实验报告各实验共用。

包含：中英文混排字体设置、代码框（Python 语法高亮）、要求/结果框、
公式（居中 + 右编号）、插图（居中 + 图注，自动编号）、封面信息表。

风格约定（按实验一最终定稿）：
  * 代码一律放进"代码框"（单元格表格 + 灰底 + 细边框）并做语法着色
  * 公式 = mathtext 图片，居中排版、右侧自动编号
  * 图注格式：图 <实验号>-<序号>　标题（不带来源/主题后缀）
  * 不写 a/b/c/d 小标题，靠【实验要求】框 + 【代码实现】/【原理与推导】/【运行结果】分段
"""
import io
import keyword
import os
import builtins
import tokenize

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

CN_FONT, CN_HEI, MONO = "宋体", "黑体", "Consolas"
TEXT_W_CM = 16.0
EQ_SCALE = 0.62
FIG_MAX_IN = 5.7


# ---------------------------------------------------------------------------
# 字体与段落
# ---------------------------------------------------------------------------
def set_font(run, size=10.5, bold=False, cn=CN_FONT, ascii_font="Times New Roman",
             color=None, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), cn)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def para(doc, text="", size=10.5, bold=False, cn=CN_FONT, ascii_font="Times New Roman",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent_chars=0, before=0, after=4,
         line=1.25, color=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if indent_chars:
        pf.first_line_indent = Pt(size * indent_chars)
    if text:
        set_font(p.add_run(text), size=size, bold=bold, cn=cn, ascii_font=ascii_font, color=color)
    return p


def h1(doc, text):
    return para(doc, text, size=14, bold=True, cn=CN_HEI, ascii_font=CN_HEI,
                align=WD_ALIGN_PARAGRAPH.LEFT, before=14, after=6, line=1.2)


def h2(doc, text):
    return para(doc, text, size=12, bold=True, cn=CN_HEI, ascii_font=CN_HEI,
                align=WD_ALIGN_PARAGRAPH.LEFT, before=10, after=5, line=1.2)


def body(doc, text):
    return para(doc, text, indent_chars=2)


def label(doc, text, color="1F4E79"):
    return para(doc, text, size=10.5, bold=True, cn=CN_HEI, ascii_font=CN_HEI,
                align=WD_ALIGN_PARAGRAPH.LEFT, before=8, after=3, color=color)


# ---------------------------------------------------------------------------
# 方框（要求框 / 代码框 / 结果框）
# ---------------------------------------------------------------------------
def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def cell_margins(table, top=60, start=100, bottom=60, end=100):
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement("w:" + tag)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    table._tbl.tblPr.append(mar)


def one_cell_box(doc, fill, border="BFBFBF"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell = t.cell(0, 0)
    cell.width = Cm(TEXT_W_CM)
    shade(cell, fill)
    cell_margins(t)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), border)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    cell.paragraphs[0].text = ""
    return cell


def py_segments(code):
    """把源码切成 (文本, 颜色) 片段；拼接结果与源码逐字符一致。"""
    try:
        toks = list(tokenize.generate_tokens(io.StringIO(code).readline))
    except Exception:
        return [(code, "1F1F1F")]

    line_start, off = [], 0
    for ln in code.split("\n"):
        line_start.append(off)
        off += len(ln) + 1

    def to_off(pos):
        r, c = pos
        if 1 <= r <= len(line_start):
            return line_start[r - 1] + c
        return len(code)

    skip = {tokenize.ENCODING, tokenize.ENDMARKER, tokenize.NEWLINE,
            tokenize.NL, tokenize.INDENT, tokenize.DEDENT}
    segs, cursor, pending_def = [], 0, False
    for t in toks:
        if t.type in skip or not t.string:
            continue
        color = "1F1F1F"
        if t.type == tokenize.COMMENT:
            color = "008000"
        elif t.type == tokenize.STRING:
            color = "A31515"
        elif t.type == tokenize.NUMBER:
            color = "098658"
        elif t.type == tokenize.NAME:
            if keyword.iskeyword(t.string):
                color = "0000FF"
                pending_def = t.string in ("def", "class")
            elif pending_def:
                color = "267F99"
                pending_def = False
            elif t.string in dir(builtins):
                color = "795E26"
        s, e = to_off(t.start), to_off(t.end)
        if s > cursor:
            segs.append((code[cursor:s], "1F1F1F"))
        segs.append((code[s:e], color))
        cursor = e
    if cursor < len(code):
        segs.append((code[cursor:], "1F1F1F"))
    return segs


def code_box(doc, code):
    cell = one_cell_box(doc, "F7F7F7")
    cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
    lines = [[]]
    for text, color in py_segments(code.rstrip("\n")):
        parts = text.split("\n")
        for k, part in enumerate(parts):
            if k:
                lines.append([])
            if part:
                lines[-1].append((part, color))
    for ln in lines:
        p = cell.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = Pt(12)
        for text, color in ln:
            set_font(p.add_run(text), size=9, cn=CN_FONT, ascii_font=MONO, color=color)
    return cell


def req_box(doc, text):
    cell = one_cell_box(doc, "EAF1F8", border="9DC3E6")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.2
    set_font(p.add_run("实验要求　"), size=10.5, bold=True, cn=CN_HEI,
             ascii_font=CN_HEI, color="1F4E79")
    set_font(p.add_run(text), size=10.5)
    return cell


def out_box(doc, text):
    cell = one_cell_box(doc, "FAFAFA", border="C9C9C9")
    cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
    for ln in text.rstrip("\n").split("\n"):
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(12)
        set_font(p.add_run(ln if ln else " "), size=9, cn=CN_FONT,
                 ascii_font=MONO, color="1F1F1F")
    return cell


# ---------------------------------------------------------------------------
# 编号、公式、插图
# ---------------------------------------------------------------------------
class Numbering:
    """图号与公式号计数器，prefix 为实验序号（'1'、'2' …）。"""

    def __init__(self, prefix="1"):
        self.prefix = prefix
        self.fig = 0
        self.eq = 0

    def next_fig(self):
        self.fig += 1
        return f"图 {self.prefix}-{self.fig}"

    def next_eq(self):
        self.eq += 1
        return f"({self.prefix}-{self.eq})"


def add_eq(doc, path, num):
    """公式：普通段落 + 居中/右对齐制表位（放入表格会被 LibreOffice 裁切）。"""
    w, h = Image.open(path).size
    label_txt = num.next_eq()
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(5)
    pf.space_after = Pt(5)
    pf.line_spacing = 1.0
    pf.tab_stops.add_tab_stop(Cm(TEXT_W_CM / 2), WD_TAB_ALIGNMENT.CENTER)
    pf.tab_stops.add_tab_stop(Cm(TEXT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
    scale = min(EQ_SCALE, 11.6 * 300 / (w * 2.54))
    run = p.add_run()
    run.add_tab()
    run.add_picture(path, width=Cm(w / 300 * 2.54 * scale))
    run.add_tab()
    set_font(p.add_run(label_txt), size=10.5)
    return label_txt


def add_fig(doc, path, caption, num, max_in=FIG_MAX_IN):
    w, h = Image.open(path).size
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(path, width=Cm(max_in * 2.54))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.line_spacing = 1.0
    label_txt = num.next_fig()
    set_font(cap.add_run(f"{label_txt}\u3000{caption}"), size=9, color="404040")
    return label_txt


# ---------------------------------------------------------------------------
# 文档骨架
# ---------------------------------------------------------------------------
def new_doc():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.left_margin = sec.right_margin = Cm(2.5)
    return doc


def cover_table(doc, date, location="B10 楼 301 房", preview="已预习", teacher="李成安"):
    rows = [
        ("地　　点：", location, "实验台号：", ""),
        ("实验日期与时间：", date, "评　　分：", ""),
        ("预习检查纪录：", preview, "实验教师：", teacher),
    ]
    t = doc.add_table(rows=3, cols=4)
    t.style = "Table Grid"
    widths = (3.4, 4.6, 3.4, 4.6)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = t.cell(r, c)
            cell.width = Cm(widths[c])
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            set_font(p.add_run(val), size=10.5, bold=(c % 2 == 0))
    return t


def doc_title(doc, text):
    return para(doc, text, size=16, bold=True, cn=CN_HEI, ascii_font=CN_HEI,
                align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=12)
