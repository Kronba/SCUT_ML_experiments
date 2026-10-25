# -*- coding: utf-8 -*-
"""实验二：调包实现 + 自实现 的新版报告，写入 总实验报告 (Copy).docx。

以 docx_kit 构建新版"实验二"章节，再整体替换原文档中的实验二章节；
实验一、三、四保持不动。
"""
import json
import os
import shutil
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_kit as K
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
REPORT_DIR = os.path.dirname(HERE)
PROJ_DIR = os.path.dirname(REPORT_DIR)
COPY = os.path.join(REPORT_DIR, "总实验报告 (Copy).docx")
ASSETS = os.path.join(REPORT_DIR, "素材")
NB = json.load(open(os.path.join(PROJ_DIR, "exp2.ipynb"), encoding="utf-8"))

NUM = K.Numbering("2")
CJK = "微软雅黑"


def code(i):
    return "".join(NB["cells"][i]["source"])


def out_text(i):
    parts = []
    for o in NB["cells"][i].get("outputs", []):
        if o.get("output_type") == "stream":
            parts.append("".join(o["text"]))
        elif o.get("output_type") == "execute_result":
            parts.append("".join(o["data"].get("text/plain", [])))
    return "".join(parts).strip("\n")


def J(*parts):
    return os.path.join(ASSETS, "exp2", *parts)


def eq(doc, key):
    return K.add_eq(doc, os.path.join(ASSETS, "exp2", "formula", key + ".png"), NUM)


def fig(doc, path, caption):
    return K.add_fig(doc, path, caption, NUM)


# ---------------------------------------------------------------------------
def msyh(doc, text, size, bold=False, align=None, indent_pt=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if indent_pt is not None:
        p.paragraph_format.first_line_indent = Pt(indent_pt)
    K.set_font(p.add_run(text), size=size, bold=bold, cn=CJK, ascii_font=CJK)
    return p


def doc_title(doc, text):
    p = msyh(doc, text, 18, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.first_line_indent = Pt(0)
    return p


def h1(doc, text):
    return msyh(doc, text, 14, bold=True)


def h2(doc, text):
    return msyh(doc, text, 12, bold=True)


def item(doc, text):
    return msyh(doc, text, 12, indent_pt=21)


def data_table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        p = t.cell(0, j).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        K.set_font(p.add_run(h), size=9.5, bold=True)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            p = t.cell(i, j).paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            K.set_font(p.add_run(str(v)), size=9.5)
    return t


def model_code(path, marker):
    text = open(path, encoding="utf-8").read()
    i = text.index(marker)
    return text[:i].rstrip(), text[i:].strip()


# ===========================================================================
def build_section(doc):
    doc_title(doc, "实验二 支持向量机及核函数的使用")

    # ---------------------------------------------------------------- 一
    h1(doc, "一、支持向量机")
    h2(doc, "1.实验目标")
    item(doc, "a.理解支持向量机的原理与应用方法")
    item(doc, "b.学习如何使用核函数将分类任务转化为线性或非线性可分的问题。")
    h2(doc, "2. 实验数据")
    item(doc, "将Auto数据集中的mpg按照中位数分为两类，新增一个变量grade，并用0和1分别表示。")
    h2(doc, "3. 实验要求")
    for s in ["a.编写Python代码实现支持向量机。",
              "b.加载数据集后，进行基本的数据清洗和处理.",
              "c.将数据集分为训练集和测试集，使用支持向量机模型进行训练，选择合适的核函数如"
              "RBF核函数。（可以尝试使用其他核函数）",
              "d.在测试集上进行预测，并计算模型的准确率、混淆矩阵、精确率等指标。",
              "e.绘制分类结果的混淆矩阵图。"]:
        item(doc, s)

    h2(doc, "4.实验步骤")
    K.body(doc, "本实验在 Python 3.12 环境下完成：pandas 负责数据读取与整理，scikit-learn 提供"
                "SVM 与感知机的调包实现，自实现则封装在 models/exp2/ 下——svm.py 用 SMO 算法"
                "求解对偶问题，perceptron.py 实现误分类驱动的在线学习，接口与 scikit-learn "
                "保持一致。首先导入依赖库与自实现模型：")
    K.code_box(doc, code(1))

    # ---- a
    K.req_box(doc, "a.编写Python代码实现支持向量机。")
    K.label(doc, "【代码实现】")
    K.body(doc, "调包实现直接创建 scikit-learn 的 SVC：")
    K.code_box(doc, code(3))
    K.body(doc, "自实现的 SVC 封装在 models/exp2/svm.py 中：用 SMO（序列最小最优化）算法求解"
                "对偶问题，支持线性核、多项式核与高斯核。文件前半部分为核函数与参数处理：")
    svm1, svm2 = model_code(os.path.join(PROJ_DIR, "models", "exp2", "svm.py"),
                            "    def fit(self, X, y):")
    K.code_box(doc, svm1)
    K.body(doc, "文件后半部分为 SMO 训练与预测接口：")
    K.code_box(doc, svm2)

    K.label(doc, "【原理与推导】")
    K.body(doc, "支持向量机的基本思想是在特征空间中寻找一个既能分开两类样本、又使间隔尽可能"
                "大的超平面，该超平面可写为：")
    e1 = eq(doc, "f2_hyperplane")
    K.body(doc, f"如式 {e1} 所示，样本到超平面的距离由 |wᵗx + b| / ‖w‖ 决定。若要求所有样本"
                "都被正确分类且位于间隔之外，即满足 yᵢ(wᵗxᵢ + b) ≥ 1，则两类之间的间隔宽度为：")
    e2 = eq(doc, "f2_margin")
    K.body(doc, f"式 {e2} 表明间隔只与权向量的模长有关，模长越小间隔越大。于是学习问题转化为"
                "一个带线性约束的凸二次规划：")
    e3 = eq(doc, "f2_hard")
    K.body(doc, f"式 {e3} 即硬间隔支持向量机。决定间隔的只有落在间隔边界上的少数样本，它们"
                "称为支持向量；删除其他样本不影响解，这正是支持向量机解的稀疏性。")
    fig(doc, J("theory", "fig2_svm_margin.jpg"), "最大间隔超平面与支持向量")
    K.body(doc, "实际数据往往线性不可分，为每个样本引入松弛变量 ξᵢ 与惩罚系数 C，允许少量样本"
                "进入间隔甚至被错分，得到软间隔支持向量机：")
    e4 = eq(doc, "f2_soft")
    K.body(doc, f"式 {e4} 中 C 控制对违反间隔的惩罚力度：C 越大，模型越不容忍错分，间隔被压缩"
                "甚至过拟合；C 越小，模型允许更多样本进入间隔，间隔更宽。软间隔目标也可以"
                "写成合页损失加 L2 正则的形式：")
    e5 = eq(doc, "f2_hinge")
    fig(doc, J("theory", "fig2_soft_margin_C.jpg"), "惩罚系数 C 对间隔宽度与容错能力的影响")
    K.body(doc, "求解上述二次规划时，引入拉格朗日乘子 αᵢ 可得到只含内积的对偶问题，这也是"
                "核方法能够生效的前提：")
    e6 = eq(doc, "f2_lagrange")
    e7 = eq(doc, "f2_dual")
    K.body(doc, f"式 {e6} 为拉格朗日函数，式 {e7} 为对偶问题。由 KKT 条件可以推出：只有 αᵢ > 0 "
                "的样本（支持向量）参与权向量的构造，且权向量是支持向量的线性组合：")
    e8 = eq(doc, "f2_kkt")
    K.body(doc, "预测时只需计算待判样本与支持向量的核函数值：")
    e9 = eq(doc, "f2_decision")
    K.body(doc, f"式 {e9} 就是自实现中 SMO 求解完成后用于推理的判别式；SMO 每次只挑选一对"
                "乘子 αᵢ、αⱼ 做解析更新，反复迭代直到满足 KKT 条件。")
    K.label(doc, "【运行结果】")
    K.body(doc, "该单元格只创建模型对象，尚未训练，因此没有输出。")

    # ---- b
    K.req_box(doc, "b.加载数据集后，进行基本的数据清洗和处理.")
    K.label(doc, "【代码实现】")
    K.code_box(doc, code(5))
    K.code_box(doc, code(6))
    K.code_box(doc, code(7))
    K.label(doc, "【原理与推导】")
    K.body(doc, "Auto MPG 数据集共 398 条样本、9 个字段（含汽车名称）。其中 horsepower 有 6 个"
                "缺失值，若直接建模会因缺失而报错或引入偏差，故先删除含有缺失的样本，处理后"
                "剩余 392 条。")
    K.body(doc, "本实验是回归目标转分类任务的典型做法：以 mpg 的中位数为阈值，大于等于中位数"
                "的样本记为 1（高油耗里程），其余记为 0。由于阈值取中位数，变换后两类样本数量"
                "完全相同（各 196 条），是一组天然平衡的二分类数据，评估时不必担心类别偏斜。")
    fig(doc, J("theory", "fig2_mpg_binning.jpg"), "按 mpg 中位数构造二分类标签")
    K.label(doc, "【运行结果】")
    K.body(doc, "horsepower 列的非空数量为 392，缺失的 6 条被识别为 NaN；构造标签后两类的"
                "样本数均为 196：")
    K.out_box(doc, out_text(5))
    K.out_box(doc, out_text(7))

    # ---- c
    K.req_box(doc, "c.将数据集分为训练集和测试集，使用支持向量机模型进行训练，选择合适的核函数"
                   "如RBF核函数。（可以尝试使用其他核函数）")
    K.label(doc, "【代码实现】")
    K.code_box(doc, code(9))
    K.code_box(doc, code(10))
    K.code_box(doc, code(11))
    K.body(doc, "自实现：创建 SVC 实例并训练，随后与调包实现对比：")
    K.code_box(doc, code(17))
    K.label(doc, "【原理与推导】")
    K.body(doc, "特征矩阵取 cylinders、displacement、horsepower、weight、acceleration、"
                "model_year、origin 共 7 个数值特征（汽车名称为文本，不参与建模），并采用"
                "分层采样按 8 : 2 划分训练集与测试集。支持向量机的间隔与核函数都建立在欧氏"
                "距离之上，特征量纲不同会使距离被量纲大的特征主导，因此先做标准化：")
    e10 = eq(doc, "f2_scale")
    K.body(doc, f"式 {e10} 的均值与标准差只在训练集上估计，再用于测试集。")
    K.body(doc, "当样本在原始空间中线性不可分时，可以通过核函数把样本隐式地映射到高维特征"
                "空间，使其在新空间中线性可分，而计算仍然只涉及原始空间的内积：")
    e11 = eq(doc, "f2_kernel_def")
    e12 = eq(doc, "f2_rbf")
    e13 = eq(doc, "f2_poly")
    K.body(doc, f"式 {e11} 是核函数的定义；式 {e12} 为高斯（RBF）核，γ 越大单个样本的影响"
                f"范围越小；式 {e13} 为多项式核。下图以同心圆数据说明核技巧的作用。")
    fig(doc, J("theory", "fig2_kernel_trick.jpg"),
        "核技巧：二维线性不可分的数据在三维空间中线性可分")
    K.body(doc, "为验证核函数的选择，在同一份数据上分别使用线性核、三次多项式核与高斯核，"
                "并比较测试集准确率（为便于在二维平面上绘制决策边界，这里只取 weight 与 "
                "horsepower 两个特征，因此准确率与使用全部 7 个特征时不同）：")
    fig(doc, J("theory", "fig2_kernel_compare.jpg"), "不同核函数的决策边界与测试集准确率对比")
    K.body(doc, "高斯核中 γ 的取值直接影响模型的复杂度：γ 过小会使决策边界过于平滑而欠拟合，"
                "γ 过大则会使每个样本各自形成一个小区域而过拟合：")
    fig(doc, J("theory", "fig2_gamma_effect.jpg"), "核参数 γ 对决策边界与拟合程度的影响")
    K.body(doc, "为进一步确定合适的参数组合，对惩罚系数 C 与核参数 γ 做 5 折交叉验证网格搜索，"
                "图中红框为最优组合：")
    fig(doc, J("theory", "fig2_param_grid.jpg"), "C 与 γ 的网格搜索结果（5 折交叉验证准确率）")
    K.label(doc, "【运行结果】")
    K.body(doc, "两种实现的结果完全一致：准确率均为 89.87%，支持向量均为 89 个，测试集的预测"
                "逐样本相同。")
    data_table(doc, ["实现方式", "准确率", "支持向量数", "预测一致率"],
               [["sklearn 调包", "89.87%", "89", "—"],
                ["自实现（SMO）", "89.87%", "89", "与调包 100% 一致"]])
    K.body(doc, "为进一步检验两者是否真的等价，把两套实现的决策边界与决策函数值放在一起"
                "比较：")
    fig(doc, J("compare", "fig2_svm_compare.jpg"), "两种 SVM 实现的决策边界与决策值对比")
    K.body(doc, "在二维特征上，两条决策边界几乎完全重合；在 7 个特征上，两套实现的决策函数值"
                "最大差异仅 1.54×10⁻³，这一微小差异来自 SMO 的收敛容差（tol = 10⁻³）"
                "与 scikit-learn 所用 libsvm 的实现细节不同。")

    # ---- d
    K.req_box(doc, "d.在测试集上进行预测，并计算模型的准确率、混淆矩阵、精确率等指标。")
    K.label(doc, "【代码实现】")
    K.code_box(doc, code(13))
    K.code_box(doc, code(14))
    K.code_box(doc, code(15))
    K.code_box(doc, code(16))
    K.label(doc, "【原理与推导】")
    K.body(doc, "分类任务的准确率定义为判对样本占样本总数的比例：")
    e14 = eq(doc, "f2_acc")
    K.body(doc, "准确率在类别不平衡时容易高估模型效果，因此还需要混淆矩阵以及由它导出的"
                "查准率、查全率与 F1 分数：")
    e15 = eq(doc, "f2_prf")
    K.body(doc, f"式 {e15} 中，查准率 P 衡量“被判为正例的样本中真正例的比例”，查全率 R 衡量"
                "“真正的正例中被找出来的比例”，F1 是二者的调和平均。综合两类表现时还可使用"
                "宏平均与加权平均：")
    e16 = eq(doc, "f2_macro")
    K.label(doc, "【运行结果】")
    K.body(doc, "在 79 个测试样本上，模型准确率为 89.87%：")
    K.out_box(doc, out_text(14))
    K.body(doc, "混淆矩阵显示：40 个低 mpg 样本中 34 个判对、6 个判错；39 个高 mpg 样本中 "
                "37 个判对、2 个判错：")
    K.out_box(doc, out_text(15))
    K.body(doc, "分类报告给出每一类的查准率、查全率与 F1 分数，宏平均与加权平均均为 0.90：")
    K.out_box(doc, out_text(16))

    # ---- e
    K.req_box(doc, "e.绘制分类结果的混淆矩阵图。")
    K.label(doc, "【代码实现】")
    K.code_box(doc, code(19))
    K.body(doc, "为便于对比，在同一张图中并排画出两种实现的混淆矩阵：")
    K.code_box(doc, code(20))
    K.label(doc, "【原理与推导】")
    K.body(doc, "混淆矩阵的可视化把四个计数放到同一张热力图上：对角线颜色越深说明判对的样本"
                "越多，非对角线上的格子则对应两类错误。把两套实现并排放在一起，可以直观看出"
                "两者的错误分布是否一致。")
    K.label(doc, "【运行结果】")
    fig(doc, J("notebook_out", "cell20_0.jpg"), "两种 SVM 实现的混淆矩阵对比")

    # ---------------------------------------------------------------- 二
    doc.add_page_break()
    h1(doc, "二、感知机")
    h2(doc, "1.实验目标")
    item(doc, "a.理解感知机的基本原理及其应用。")
    item(doc, "b.掌握感知机模型的构建和评估方法。")
    h2(doc, "2. 实验数据")
    item(doc, "使用糖尿病数据集(diabetes.csv)，该数据集包含多个特征，用于预测患者是否患有"
              "糖尿病。")
    h2(doc, "3. 实验要求")
    for s in ["a.编写Python代码实现感知机模型。",
              "b.实现数据的预处理，使用感知机模型进行糖尿病数据的分类。",
              "c.输出模型的准确率、混淆矩阵。",
              "d.可视化特征权重，以分析模型性能和特征重要性。"]:
        item(doc, s)

    h2(doc, "4.实验步骤")
    K.req_box(doc, "a.编写Python代码实现感知机模型。")
    K.label(doc, "【代码实现】")
    K.body(doc, "调包实现创建 scikit-learn 的 Perceptron：")
    K.code_box(doc, code(24))
    K.body(doc, "自实现的感知机封装在 models/exp2/perceptron.py 中：用误分类驱动的在线学习"
                "逐样本更新参数，接口与 scikit-learn 一致：")
    perc_code = open(os.path.join(PROJ_DIR, "models", "exp2", "perceptron.py"), encoding="utf-8").read()
    K.code_box(doc, perc_code)

    K.label(doc, "【原理与推导】")
    K.body(doc, "感知机是最早的线性分类模型，它直接对样本输出类别而无需概率解释：")
    e17 = eq(doc, "f2_perc_model")
    K.body(doc, f"式 {e17} 中 sign(·) 为符号函数。感知机采用误分类驱动的学习策略：只统计被"
                "当前模型判错的样本集合 M，以误分类样本到超平面的距离之和（略去 1/‖w‖ 的"
                "尺度因子）作为损失函数：")
    e18 = eq(doc, "f2_perc_loss")
    K.body(doc, f"式 {e18} 对 w 与 b 求梯度：")
    e19 = eq(doc, "f2_perc_grad")
    K.body(doc, f"式 {e19} 给出随机梯度下降的更新规则，即每发现一个误分类样本就沿其方向"
                "调整一次参数：")
    e20 = eq(doc, "f2_perc_update")
    K.body(doc, "代码中 Perceptron(max_iter=1000, tol=1e-3, random_state=42) 表示最多迭代 "
                "1000 轮，若误分类数的下降幅度连续 5 轮小于 1e-3 则提前停止。下图给出决策"
                "边界随迭代的变化以及训练错分率曲线：")
    fig(doc, J("theory", "fig2_perceptron_update.jpg"),
        "感知机决策边界的迭代过程与训练错分率变化")
    K.body(doc, "Novikoff 定理指出：若训练集线性可分，感知机一定能在有限步内收敛，且迭代"
                "次数存在上界：")
    e21 = eq(doc, "f2_novikoff")
    K.body(doc, f"式 {e21} 中 γ 为样本到最优超平面的最小函数间隔，R 为样本的最大模长。需要"
                "注意的是，感知机只要求把样本分开，满足条件的解并不唯一，它也不会像支持向量"
                "机那样最大化间隔。下图对比了两种模型在同一数据上的决策边界。")
    fig(doc, J("theory", "fig2_perceptron_vs_svm.jpg"),
        "感知机与线性 SVM 的决策边界对比（仅取两个特征以便可视化）")
    K.label(doc, "【运行结果】")
    K.body(doc, "该单元格只创建感知机对象，尚未训练，因此没有输出。")

    K.req_box(doc, "b.实现数据的预处理，使用感知机模型进行糖尿病数据的分类。")
    K.label(doc, "【代码实现】")
    K.code_box(doc, code(22))
    K.code_box(doc, code(26))
    K.code_box(doc, code(27))
    K.code_box(doc, code(28))
    K.code_box(doc, code(29))
    K.label(doc, "【原理与推导】")
    K.body(doc, "糖尿病数据集共 768 条样本、8 个数值特征（怀孕次数、血糖、血压、皮褶厚度、"
                "胰岛素、BMI、糖尿病家族史函数、年龄）与 1 个标签 Outcome，其中未患病 500 例、"
                "患病 268 例。若模型不加学习、一律预测为未患病，准确率也能达到：")
    e22 = eq(doc, "f2_baseline")
    K.body(doc, f"式 {e22} 即多数类基线，该数据上为 65.1%，是评价模型是否真正学到信息的"
                "重要参照。")
    fig(doc, J("theory", "fig2_class_balance.jpg"), "糖尿病数据集的类别分布与多数类基线")
    K.body(doc, "数据表中没有显式的缺失值，但 Glucose、BloodPressure、SkinThickness、Insulin、"
                "BMI 等字段存在取值为 0 的记录（如 SkinThickness 有 227 条、Insulin 有 374 条），"
                "这些取值在生理上不可能出现，应视为缺失。本实验按 notebook 的处理方式保留原始"
                "取值，未做插补，这也是模型性能受限的原因之一。")
    K.body(doc, f"训练前对特征做标准化（同式 {e10}），随后按 8 : 2 分层划分训练集与测试集。")
    K.label(doc, "【运行结果】")
    K.body(doc, "读取后打印数据结构：共 768 条样本、9 个字段，无缺失值；模型拟合完成后的"
                "参数为：")
    K.out_box(doc, out_text(22))
    K.out_box(doc, out_text(28))

    K.req_box(doc, "c.输出模型的准确率、混淆矩阵。")
    K.label(doc, "【代码实现】")
    K.code_box(doc, code(31))
    K.code_box(doc, code(32))
    K.code_box(doc, code(33))
    K.body(doc, "自实现：创建 Perceptron 实例并训练，随后与调包实现对比：")
    K.code_box(doc, code(34))
    K.label(doc, "【原理与推导】")
    K.body(doc, f"评价指标与支持向量机部分相同，仍使用准确率（式 {e14}）、混淆矩阵以及查准率、"
                f"查全率与 F1 分数（式 {e15}）。由于该数据存在类别不平衡，评估时需要把结果"
                f"与多数类基线（式 {e22}，65.1%）对照来看。")
    K.label(doc, "【运行结果】")
    K.body(doc, "调包实现的准确率为 72.73%，自实现为 68.83%，两者相差 3.9 个百分点：")
    K.out_box(doc, out_text(31))
    K.out_box(doc, out_text(32))
    K.out_box(doc, out_text(33))
    K.body(doc, "自实现的对比结果如下：两种实现的预测一致率为 79.22%，权重相关系数为 0.6633。")
    K.out_box(doc, out_text(34))
    K.body(doc, "两者的差异并不是实现错误，而是感知机本身的性质决定的：训练集线性不可分时，"
                "感知机的解不唯一，最终参数取决于样本的更新顺序。把自实现在 11 个不同随机种子下"
                "重复训练，测试准确率在 63.64%~74.03% 之间波动（均值 68.60%），sklearn 的 "
                "72.73% 正落在这一区间内：")
    fig(doc, J("compare", "fig2_perceptron_compare.jpg"),
        "两种感知机实现的决策边界与不同更新顺序下的准确率")
    K.body(doc, "这一现象恰好印证了感知机与支持向量机的本质区别：感知机只要求把样本分开，"
                "而支持向量机在所有可行解中挑出间隔最大的那个，因此后者的解唯一且更稳定。")

    K.req_box(doc, "d.可视化特征权重，以分析模型性能和特征重要性。")
    K.label(doc, "【代码实现】")
    K.code_box(doc, code(37))
    K.code_box(doc, code(38))
    K.label(doc, "【原理与推导】")
    K.body(doc, f"由式 {e17} 可知，特征权重 w 决定了每个特征对判决的方向与强度：权重为正表示"
                "该特征取值越大越倾向于判为患病，为负则相反；绝对值越大说明该特征对分类的影响"
                "越强。由于训练前所有特征都已标准化，不同特征之间的权重具有可比性，因此可以"
                "直接用权重的绝对值衡量特征重要性。")
    K.body(doc, "从结果看，Glucose（血糖）的权重最大，与临床上血糖是糖尿病主要判别指标的事实"
                "一致；BMI 与 Age 也给出正的贡献；而 SkinThickness、Insulin 等特征的权重接近 0，"
                "说明在线性模型下它们对判决的贡献有限——这与 b 中提到的这些字段存在大量 0 值"
                "（相当于缺失）有关。")
    K.label(doc, "【运行结果】")
    fig(doc, J("notebook_out", "cell38_0.jpg"), "感知机各特征的权重（按绝对值升序排列）")
    K.body(doc, "把两种实现的权重向量放在同一坐标系中对比，可以看到整体趋势一致（相关系数 "
                "0.6633），但少数特征的权重差异较大——这正是更新顺序不同导致的两个不同解：")
    fig(doc, J("notebook_out", "cell35_1.jpg"), "两种感知机实现的权重对比")

    doc.add_page_break()


def main():
    backup = "/tmp/总实验报告_Copy_实验一后备份.docx"
    shutil.copy(COPY, backup)
    print("已备份 ->", backup)

    doc = Document(COPY)
    body = doc.element.body

    stop = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("实验三"):
            stop = p._element
            break
    if stop is None:
        raise RuntimeError("未找到实验三标题")

    start = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("实验二"):
            start = p._element
            break
    if start is None:
        raise RuntimeError("未找到实验二标题")

    # 保留实验二的封面信息表（标题后的第一张表）
    cover = None
    el = start
    while True:
        el = el.getnext()
        if el is None or el is stop:
            break
        if el.tag == qn("w:tbl"):
            cover = el
            break

    # 在文档末尾构建新章节
    sect = body.find(qn("w:sectPr"))
    start_len = list(body).index(sect) if sect is not None else len(list(body))
    build_section(doc)
    new_els = [el for el in list(body)[start_len:] if el.tag != qn("w:sectPr")]
    print(f"新章节元素数：{len(new_els)}")

    # 删除旧章节（从实验二标题到实验三标题之前）
    removed = 0
    el = start
    while el is not None and el is not stop:
        nxt = el.getnext()
        if el is not cover:
            body.remove(el)
            removed += 1
        el = nxt
    print(f"已删除旧章节元素：{removed}")

    # 插入到实验三之前
    for el in new_els:
        stop.addprevious(el)

    # 封面表移到新标题之后
    if cover is not None:
        new_els[0].addnext(cover)

    doc.save(COPY)
    print("已写入 ->", COPY)
    print(f"插图 {NUM.fig} 张，公式 {NUM.eq} 个")


if __name__ == "__main__":
    main()
