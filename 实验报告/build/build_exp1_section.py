# -*- coding: utf-8 -*-
"""实验一：调包实现 + 自实现 的新版报告，写入 总实验报告 (Copy).docx。

做法：以 docx_kit 构建新版"实验一"章节（含两套实现的代码与结果对比），
再用底层 XML 操作把原文档中的实验一章节整体替换掉，其余实验保持不动。
"""
import json
import os
import shutil
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import docx_kit as K
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
REPORT_DIR = os.path.dirname(HERE)
PROJ_DIR = os.path.dirname(REPORT_DIR)
COPY = os.path.join(REPORT_DIR, "总实验报告 (Copy).docx")
ASSETS = os.path.join(REPORT_DIR, "素材")
NB = json.load(open(os.path.join(PROJ_DIR, "exp1.ipynb"), encoding="utf-8"))

NUM = K.Numbering("1")
CJK = "微软雅黑"          # 汇总本中标题/目标/要求等使用微软雅黑


def code(i):
    return "".join(NB["cells"][i]["source"])


def out_text(i):
    parts = []
    for o in NB["cells"][i].get("outputs", []):
        if o.get("output_type") == "stream":
            parts.append("".join(o["text"]))
        elif o.get("output_type") == "execute_result":
            parts.append("".join(o["data"].get("text/plain", [])))
    return "".join(parts).strip("\n")


def J(*parts):
    return os.path.join(ASSETS, *parts)


def eq(doc, key):
    return K.add_eq(doc, J("formula", key + ".png"), NUM)


def fig(doc, path, caption):
    return K.add_fig(doc, path, caption, NUM)


# ---------------------------------------------------------------------------
# 汇总本使用的标题样式
# ---------------------------------------------------------------------------
def msyh(doc, text, size, bold=False, align=None, indent_pt=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if indent_pt is not None:
        p.paragraph_format.first_line_indent = Pt(indent_pt)
    K.set_font(p.add_run(text), size=size, bold=bold, cn=CJK, ascii_font=CJK)
    return p


def doc_title(doc, text):
    p = msyh(doc, text, 18, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.first_line_indent = Pt(0)
    return p


def h1(doc, text):
    return msyh(doc, text, 14, bold=True)


def h2(doc, text):
    return msyh(doc, text, 12, bold=True)


def item(doc, text):
    return msyh(doc, text, 12, indent_pt=21)


def data_table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        p = t.cell(0, j).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        K.set_font(p.add_run(h), size=9.5, bold=True)
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            p = t.cell(i, j).paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            K.set_font(p.add_run(str(v)), size=9.5)
    return t


def model_code(path, marker):
    text = open(path, encoding="utf-8").read()
    i = text.index(marker)
    return text[:i].rstrip(), text[i:].strip()


# ===========================================================================
def build_section(doc):
    """把新版实验一章节追加到 doc 末尾。"""
    doc_title(doc, "实验一 逻辑斯蒂回归 线性判别分析")

    # ---------------------------------------------------------------- 一
    h1(doc, "一、逻辑斯蒂回归")
    h2(doc, "1. 实验目标")
    item(doc, "理解并实现Logistic 回归算法")
    item(doc, "掌握如何使用逻辑回归解决实际问题。")
    h2(doc, "2、实验数据")
    item(doc, "Simmons 数据集(Simmons.csv) 包含50名拥有Simmons信用卡的顾客和50名没有Simmons"
              "信用卡的其他顾客组成。目标变量是顾客是否使用了优惠券。")
    h2(doc, "3、实验要求")
    for s in ["a.编写Python代码实现Logistic 回归，并简单展示Sigmoid函数。",
              "b.读取数据集，并画出散点图。",
              "c.训练Logistic回归模型，并输出最终的分类准确率。",
              "d.绘制Logistic回归模型的决策边界。"]:
        item(doc, s)

    h2(doc, "4、实验步骤")
    K.body(doc, "本实验在 Python 3.12 环境下完成：pandas 负责数据读取与整理，numpy 完成数值"
                "运算，matplotlib 负责绘图。为便于对比，本实验给出两套实现——调包实现直接调用 "
                "scikit-learn 的 LogisticRegression 与 LinearDiscriminantAnalysis；自实现则"
                "封装在 models/exp1/ 下，sigmoid.py 用 numpy 实现逻辑斯蒂回归，LDA.py 实现"
                "线性判别分析，两者的接口都与 scikit-learn 保持一致，因此可以在同一段实验"
                "代码里直接替换。首先导入依赖库与自实现模型：")
    K.code_box(doc, code(1))

    # ---- a
    K.req_box(doc, "a.编写Python代码实现Logistic 回归，并简单展示Sigmoid函数。")
    K.label(doc, "【代码实现】")
    K.body(doc, "Sigmoid 函数是逻辑斯蒂回归的核心，两种实现都基于同一公式。先给出用于展示的"
                "函数定义与绘图代码：")
    K.code_box(doc, code(3))
    K.code_box(doc, code(4))
    K.body(doc, "自实现的模型封装在 models/exp1/sigmoid.py 中：文件中包含一个数值稳定的 "
                "Sigmoid 函数与一个 LogisticRegression 类，类内提供牛顿法与梯度下降两种求解"
                "方式，并保留了与 scikit-learn 一致的接口。文件前半部分为：")
    part1, part2 = model_code(os.path.join(PROJ_DIR, "models", "exp1", "sigmoid.py"),
                              "    # ------------------------------------------------------------------\n"
                              "    # 训练与预测")
    K.code_box(doc, part1)
    K.body(doc, "文件后半部分为训练与预测接口：")
    K.code_box(doc, part2)

    K.label(doc, "【原理与推导】")
    K.body(doc, "Logistic 回归解决的是二分类问题。设样本为 d 维特征向量 x，其类别标记为 "
                "y ∈ {0, 1}，模型并不直接拟合类别，而是先构造特征的线性组合：")
    e1 = eq(doc, "f_linear")
    K.body(doc, f"如式 {e1} 所示，z 的取值覆盖整个实数轴，无法直接解释为概率。为此引入 "
                "Sigmoid 函数，把实数轴单调压缩到 (0, 1)：")
    e2 = eq(doc, "f_sigmoid")
    K.body(doc, f"式 {e2} 中 σ(z) 被解释为样本属于正类的后验概率 P(y = 1 | x)。由式 {e2} 反解"
                "即可得到 Logistic 回归的名称来源——对数几率是输入的线性函数：")
    e3 = eq(doc, "f_logit")
    K.body(doc, f"式 {e3} 表明：若将正例判定的后验概率之比取对数，则它与特征呈线性关系，"
                "这正是 Logistic 回归被称为“对数几率回归”的原因。Sigmoid 函数还有一个在参数"
                "估计中极为关键的性质，即其导数可用自身表示：")
    e4 = eq(doc, "f_sigmoid_deriv")
    K.body(doc, f"式 {e4} 是后续梯度求解能够写成简洁形式的基础。")
    K.body(doc, "自实现中的 sigmoid 采用分段写法：当 z ≥ 0 时计算 1/(1+e^{-z})，当 z < 0 时"
                "计算 e^{z}/(1+e^{z})，从而避免 exp 溢出，这也是它与直接按公式书写在数值上的"
                "唯一差别。")

    K.label(doc, "【运行结果】")
    K.body(doc, "把自实现的 Sigmoid 与直接按公式计算的版本放在一起比较：")
    fig(doc, J("exp1", "compare", "fig1_sigmoid_compare.jpg"),
        "Sigmoid 函数两种实现的对比与数值误差")
    K.body(doc, "两种实现的最大绝对误差为 5.55×10⁻¹⁷，已达到双精度浮点数的舍入误差量级，"
                "说明自实现的 Sigmoid 与公式完全一致。")

    # ---- b
    K.req_box(doc, "b.读取数据集，并画出散点图。")
    K.label(doc, "【代码实现】")
    K.code_box(doc, code(7))
    K.body(doc, "读取数据后打印数据结构信息，再以 Spending 为横轴、Card 为纵轴绘制散点图，"
                "并按是否使用优惠券（Coupon）区分颜色：")
    K.code_box(doc, code(8))
    K.label(doc, "【原理与推导】")
    K.body(doc, "Simmons 数据集共 100 个样本，含 4 个字段：Customer（顾客编号）、"
                "Spending（消费额）、Card（信用卡持有情况）、Coupon（是否使用优惠券，为"
                "目标变量）。散点图的意义在于先在二维特征空间中观察两类样本的可分性：若两类"
                "样本能够被一条直线大致分开，则线性模型是合适的选择；若严重重叠，则分类准确率"
                "的上限将受到数据本身可分性的限制。")
    K.body(doc, "在划分训练集与测试集时，代码使用留出法按 8 : 2 划分，并通过 stratify 参数"
                "实施分层采样，使训练集与测试集中正负样本的比例与原始数据保持一致。留出法只做"
                "一次划分，而 k 折交叉验证会轮流留出一折作测试集并取 k 个结果的平均，二者的"
                "原理与分层采样的效果如下图：")
    fig(doc, J("concept", "fig_cross_validation.jpg"), "评估方法：10 折交叉验证与分层留出法")
    K.label(doc, "【运行结果】")
    K.body(doc, "运行 data.info() 得到数据的结构与类型信息：")
    K.out_box(doc, out_text(7))
    fig(doc, J("exp1", "notebook_out", "cell08_0.jpg"), "Simmons 数据集散点图")
    K.body(doc, "散点图显示：Card 取值为 0/1 的两行样本在 Spending 方向上高度重叠，两类样本"
                "并不存在一条能够完全分开它们的直线，这预示了后续分类准确率不会很高。")

    # ---- c
    K.req_box(doc, "c.训练Logistic回归模型，并输出最终的分类准确率。")
    K.label(doc, "【代码实现】")
    K.body(doc, "调包实现：")
    K.code_box(doc, code(11))
    K.code_box(doc, code(12))
    K.code_box(doc, code(13))
    K.body(doc, "自实现：在 notebook 中直接创建模型实例并训练，牛顿法与梯度下降各训练一次，"
                "以便对比两种求解方式的差异：")
    K.code_box(doc, code(14))
    K.label(doc, "【原理与推导】")
    K.body(doc, "Logistic 回归的参数 w、b 采用极大似然估计求解。将 σ(z_i) 记作样本 i 属于"
                "正类的预测概率 p̂_i，则整个训练集的似然函数为：")
    e5 = eq(doc, "f_likelihood")
    K.body(doc, f"最大化式 {e5} 等价于最小化其相反数，即对数损失（log-loss），它是 Logistic "
                "回归的优化目标：")
    e6 = eq(doc, "f_logloss")
    K.body(doc, f"式 {e6} 是一个关于 w、b 的凸函数，因此不存在局部最优解的困扰。结合式 {e4}"
                f"（Sigmoid 的导数性质）对式 {e6} 求偏导，可得极为简洁的梯度形式：")
    e7 = eq(doc, "f_grad")
    K.body(doc, f"式 {e7} 表明梯度只与“预测概率与真实标记之差”有关，与线性回归的残差形式"
                "完全一致。由于目标函数二阶可导，除梯度下降外还可使用收敛更快的牛顿法：")
    e8 = eq(doc, "f_newton")
    K.body(doc, f"式 {e8} 中的 H 为海森矩阵；scikit-learn 默认采用 L-BFGS 等拟牛顿方法求解。"
                f"自实现的类把目标函数写成 0.5‖w‖² + C·Σlog-loss，与 scikit-learn 的正则化"
                "约定一致（截距不参与正则），牛顿法使用解析梯度与海森矩阵，梯度下降使用同样的"
                f"梯度并取固定步长。二者与调包实现优化的是同一个目标，因此应收敛到同一解。")
    e9 = eq(doc, "f_acc")
    K.body(doc, f"模型训练完成后，用测试集上的分类准确率衡量性能，错误率与精度的定义为式 "
                f"{e9}。")

    K.label(doc, "【运行结果】")
    K.body(doc, "三种实现的准确率完全一致，均为 60.00%（测试集 20 个样本，错分 8 个），"
                "系数与截距也几乎相同：")
    data_table(doc, ["实现方式", "准确率", "w（Spending）", "w（Card）", "截距 b"],
               [["sklearn 调包", "60.00%", "0.4227", "1.2650", "−2.5224"],
                ["自实现（牛顿法）", "60.00%", "0.4226", "1.2650", "−2.5224"],
                ["自实现（梯度下降）", "60.00%", "0.4226", "1.2650", "−2.5224"]])
    K.body(doc, "两种自实现与调包实现的系数最大差异仅为 1.4×10⁻⁵，可以认为求解结果相同；"
                "差异来源是 scikit-learn 的收敛阈值（默认 1e-4）比自实现（1e-8）更宽松。")
    K.body(doc, "虽然最终的解一致，两种求解方式的收敛速度差别很大：")
    fig(doc, J("exp1", "compare", "fig1_convergence.jpg"),
        "自实现的收敛过程：牛顿法与梯度下降")
    K.body(doc, "牛顿法只用了 4 轮迭代就收敛，梯度下降用了 1432 轮。原因是牛顿法利用了二阶"
                "导数信息，在接近最优解时具有二次收敛速度；而梯度下降只用一阶信息，且步长固定，"
                "在参数尺度差异较大时需要大量迭代。")

    # ---- d
    K.req_box(doc, "d.绘制Logistic回归模型的决策边界。")
    K.label(doc, "【代码实现】")
    K.body(doc, "调包实现的决策边界：")
    K.code_box(doc, code(17))
    K.code_box(doc, code(18))
    K.body(doc, "自实现的决策边界（用牛顿法求得的参数）：")
    K.code_box(doc, code(19))
    K.label(doc, "【原理与推导】")
    K.body(doc, "由式 " + e2.strip("()") + " 可知 σ(z) 关于 z = 0 处取值为 0.5。因此模型判定为"
                "正类与负类的分界恰好发生在 z = 0 处，即决策边界满足：")
    e10 = eq(doc, "f_boundary")
    K.body(doc, f"式 {e10} 说明：Logistic 回归的决策边界是由权向量 w 与偏置 b 确定的超平面，"
                "在二维特征空间中它退化为一条直线，w 是这条直线的法向量，其方向决定边界的"
                "朝向，b 则决定边界相对原点的偏移。代码中正是按 z = β₀ + β₁·Spending + "
                "β₂·Card = 0 反解出 Card 关于 Spending 的直线方程。")
    K.body(doc, "为进一步说明该边界的几何意义与概率解释，下图给出决策边界及两侧的概率区域：")
    fig(doc, J("theory", "fig_boundary_geo.jpg"),
        "Logistic 回归决策边界的几何解释与概率区域")
    K.body(doc, "值得注意的是，决策边界默认取在后验概率为 0.5 处，这隐含了“两类错误代价相同”"
                "的假设。若两类错误的代价不同，可通过移动分类阈值来调整边界，其综合评价手段"
                "即为 P-R 曲线与 ROC 曲线：ROC 曲线以假正例率为横轴、真正例率为纵轴，曲线下"
                "面积（AUC）可用来比较不同模型或不同阈值下的排序能力：")
    fig(doc, J("concept", "fig_roc_auc.jpg"), "ROC 曲线与 AUC")
    K.label(doc, "【运行结果】")
    K.body(doc, "把两套实现的决策边界画在同一张散点图上，并比较两者在测试集上输出的正类概率：")
    fig(doc, J("exp1", "compare", "fig1_boundary_compare.jpg"),
        "决策边界与预测概率的对比")
    K.body(doc, "两条决策边界的最大差异为 3.37×10⁻⁵，测试集预测概率的最大差异为 9.89×10⁻⁶，"
                "在图上完全重合。这说明自实现与调包实现得到的不仅是同一个准确率，而是同一条"
                "决策边界。")

    # ---------------------------------------------------------------- 二
    doc.add_page_break()
    h1(doc, "二、线性判别分析")
    h2(doc, "1.实验目标")
    item(doc, "a.理解线性判别分析(LDA)的原理及其应用。")
    item(doc, "b.学习如何使用LDA进行分类，并比较其与其他分类算法的性能。")
    h2(doc, "2.实验数据")
    item(doc, "本实验使用鸢尾花 IRIS 数据集（直接加载使用 sklearn 库中的 datasets 引用鸢尾花"
              "数据集），该数据集包含 150 个样本，分为 3 种鸢尾花，每种花的特征包含萼片长度、"
              "萼片宽度、花瓣长度和花瓣宽度。")
    h2(doc, "3.实验要求")
    for s in ["a.编写Python代码实现线性判别分析。",
              "b.加载鸢尾花数据集，并对数据进行预处理和可视化。",
              "c.使用LDA进行分类，并输出模型的准确率、混淆均值等评估指标。",
              "d.将LDA投影结果可视化，观察不同类型的分布情况。"]:
        item(doc, s)

    h2(doc, "4、实验步骤")
    K.req_box(doc, "a.编写Python代码实现线性判别分析。")
    K.label(doc, "【代码实现】")
    K.body(doc, "调包实现与自实现分别创建模型实例：")
    K.code_box(doc, code(22))
    K.body(doc, "自实现的模型封装在 models/exp1/LDA.py 中：它用 numpy 计算类均值、先验概率、"
                "类内散度矩阵 Sw 与类间散度矩阵 Sb，求解广义特征值问题得到 Fisher 投影方向，"
                "并用高斯判别函数完成分类。文件前半部分为：")
    lda1, lda2 = model_code(os.path.join(PROJ_DIR, "models", "exp1", "LDA.py"),
                            "    # ------------------------------------------------------------------\n"
                            "    # 投影与判别")
    K.code_box(doc, lda1)
    K.body(doc, "文件后半部分为投影与判别接口：")
    K.code_box(doc, lda2)

    K.label(doc, "【原理与推导】")
    K.body(doc, "线性判别分析的核心思想是“投影”：把高维样本投影到一条直线 y = wᵗx 上，"
                "并使投影结果满足——同类样本的投影点尽可能聚集，异类样本的投影点尽可能远离。"
                "为定量描述这一目标，先定义各类的均值向量与全体样本均值：")
    e11 = eq(doc, "f_class_mean")
    K.body(doc, "进而定义类内散度矩阵 S_w（衡量同类样本的分散程度）与类间散度矩阵 S_b"
                "（衡量异类均值之间的距离）：")
    e12 = eq(doc, "f_sw")
    e13 = eq(doc, "f_sb")
    K.body(doc, f"其中式 {e12} 为各类协方差矩阵之和，式 {e13} 为类均值相对全体均值的加权"
                "散布。投影后，同类样本的方差为 wᵗS_w w，异类均值之差为 wᵗ(μ₀ − μ₁)。"
                "于是 Fisher 提出用二者的比值作为优化准则：")
    e14 = eq(doc, "f_fisher")
    K.body(doc, f"式 {e14} 即 Fisher 准则，它是一个广义瑞利商。令 J(w) 对 w 求导并令其为零，"
                "可得最优投影方向所满足的方程，推导过程为：")
    e15 = eq(doc, "f_fisher_proof")
    K.body(doc, f"注意到式 {e15} 中 S_b w 的方向恒为 μ₀ − μ₁，代入并令 λ = wᵗS_b w（它只影响 "
                "w 的长度而不影响方向），即得闭式解：")
    e16 = eq(doc, "f_fisher_sol")
    K.body(doc, f"式 {e16} 表明：最优投影方向只取决于两类均值之差经类内散度矩阵“白化”之后"
                "的朝向。自实现正是按这一思路求解：先对 S_w 做 Cholesky 分解并白化，把广义"
                "特征值问题转成对称矩阵的特征值问题，再用 numpy 的 eigh 求解。")
    fig(doc, J("theory", "fig_lda_geo.jpg"), "Fisher 判别准则的几何解释：类内散度椭圆与最优投影方向")
    K.label(doc, "【运行结果】")
    K.body(doc, "该单元格只创建两个模型对象，尚未训练，因此没有输出。")

    K.req_box(doc, "b.加载鸢尾花数据集，并对数据进行预处理和可视化。")
    K.label(doc, "【代码实现】")
    K.code_box(doc, code(24))
    K.body(doc, "对特征进行标准化处理（标准化参数仅在训练集上估计）：")
    K.code_box(doc, code(25))
    K.body(doc, "绘制花瓣长度与花瓣宽度的二维散点图，观察三类的分布：")
    K.code_box(doc, code(26))
    K.label(doc, "【原理与推导】")
    K.body(doc, "鸢尾花数据集共 150 个样本、4 个连续特征、3 个类别，每类 50 个样本。代码按 "
                "8 : 2 划分训练集与测试集，并同样启用 stratify 分层采样，保证 120 个训练样本"
                "与 30 个测试样本中三类比例均为 1 : 1 : 1。")
    K.body(doc, "标准化（Z-score 规范化）将各特征变换为均值 0、方差 1 的分布：")
    e17 = eq(doc, "f_scale")
    K.body(doc, f"式 {e17} 中 μ、σ 分别为训练集上该特征的均值与标准差。对 LDA 而言，标准化"
                "并非可选步骤：求解投影方向需要计算 S_w 的逆矩阵，若各特征量纲相差悬殊，"
                "S_w 的条件数会急剧变差，投影方向将被量纲较大的特征主导。标准化后各特征处于"
                "同一尺度，Fisher 准则才能公平地权衡所有特征。")
    K.label(doc, "【运行结果】")
    fig(doc, J("exp1", "notebook_out", "cell26_0.jpg"), "鸢尾花数据集（花瓣长度—花瓣宽度）散点图")
    K.body(doc, "从图中可见：山鸢尾与另外两类在“花瓣长度”方向上几乎完全分离，而变色鸢尾与"
                "维吉尼亚鸢尾在所选的两个特征上存在少量重叠，这为下文“三分类是否能被完全正确"
                "区分”提供了直观预期。")

    K.req_box(doc, "c.使用LDA进行分类，并输出模型的准确率、混淆均值等评估指标。")
    K.label(doc, "【代码实现】")
    K.body(doc, "调包实现：")
    K.code_box(doc, code(28))
    K.code_box(doc, code(29))
    K.code_box(doc, code(30))
    K.body(doc, "自实现：训练、预测，并与调包实现的结果对比：")
    K.code_box(doc, code(31))
    K.label(doc, "【原理与推导】")
    K.body(doc, "LDA 在完成投影方向的求解后，对样本的判别方式是：先假设各类服从协方差相同的"
                "高斯分布，再由贝叶斯公式写出后验概率并取对数，得到第 k 类的判别函数：")
    e18 = eq(doc, "f_lda_disc")
    K.body(doc, f"式 {e18} 中 π_k 为先验概率，Σ 为各类共享的协方差矩阵。样本被划入判别函数"
                "取值最大的那一类。由于各类判别函数关于 x 都是线性的，任意两类之间的决策面"
                "依然是一个超平面：")
    e19 = eq(doc, "f_lda_binary_disc")
    K.body(doc, f"式 {e19} 说明了 LDA 与 Logistic 回归在形式上的统一性：二者都产生线性决策"
                "边界，区别在于参数估计的准则不同。自实现即按式 " + e18.strip("()") +
                " 计算各类判别函数并取最大者，sklearn 的默认 SVD 解法与之等价。")
    K.body(doc, "对于多分类问题，LDA 通过求解广义特征值问题得到多个投影方向，最多保留 "
                "K − 1 个（K 为类别数），从而实现有监督降维：")
    e20 = eq(doc, "f_lda_multi")
    K.body(doc, f"式 {e20} 解释了 d 中投影结果维度的由来：3 个类别最多保留 2 个判别方向。")
    K.body(doc, "在评估环节，除准确率外还输出混淆矩阵、查准率、查全率与 F1 分数。"
                "这些指标由混淆矩阵中的 TP、FP、FN、TN 四个计数导出：")
    e21 = eq(doc, "f_prf")
    K.body(doc, f"式 {e21} 中：查准率 P 衡量“被判为正例的样本中真正例的比例”，查全率 R 衡量"
                "“真正的正例中被找出来的比例”，F1 为二者的调和平均。混淆矩阵与查准率、"
                "查全率的定义如下图：")
    fig(doc, J("concept", "fig_confusion_matrix.jpg"), "混淆矩阵的构成与由它导出的查准率、查全率、F1")
    K.body(doc, "当模型输出概率并需要比较不同阈值下的表现时，可进一步使用 P-R 曲线。"
                "P-R 曲线的平衡点常用于度量曲线交叉时分类器的性能高低：")
    fig(doc, J("concept", "fig_pr_curve.jpg"), "P-R 曲线与平衡点")
    K.label(doc, "【运行结果】")
    K.body(doc, "标准化后的特征使两种实现都在测试集上取得了 100% 的分类准确率，且两者的"
                "预测完全一致：")
    data_table(doc, ["实现方式", "准确率", "类别均值数", "判别方向数", "预测一致"],
               [["sklearn 调包", "100.00%", "3", "2", "—"],
                ["自实现", "100.00%", "3", "2", "与调包完全一致"]])
    fig(doc, J("exp1", "compare", "fig1_lda_cm_compare.jpg"), "两种实现的 LDA 混淆矩阵对比")
    K.body(doc, "两个混淆矩阵都是对角阵（对角元素均为 10），说明 30 个测试样本全部判对；"
                "自实现的预测结果与调用 sklearn 得到的结果逐样本相同。")
    K.body(doc, "调包实现的分类报告给出每一类的查准率、查全率与 F1 分数，三类均为 1.00，"
                "宏平均与加权平均也均为 1.00：")
    K.out_box(doc, out_text(30))

    K.req_box(doc, "d.将LDA投影结果可视化，观察不同类型的分布情况。")
    K.label(doc, "【代码实现】")
    K.body(doc, "调包实现的投影：")
    K.code_box(doc, code(33))
    K.code_box(doc, code(34))
    K.body(doc, "自实现的投影，并与调包实现的投影作对比（投影方向可能相差一个符号，"
                "这里先统一符号再比较）：")
    K.code_box(doc, code(35))
    K.label(doc, "【原理与推导】")
    K.body(doc, f"投影变换即 y = wᵗx，其方向由式 {e16} 确定；对多分类问题则取广义特征值问题"
                f"的前 K − 1 个特征向量构成投影矩阵 V，如式 {e20}。投影后的坐标是原始 4 维"
                "特征在“类间散度最大、类内散度最小”意义下的最优二维表示，因此投影图中同类"
                "样本应当彼此聚集、不同类样本应当彼此分离。检验投影图是否满足这一形态，"
                "等价于对 Fisher 准则优化目标的一次直观验证。")
    K.label(doc, "【运行结果】")
    K.body(doc, "两种实现的投影结果相关系数均为 1.0，投影图的结构完全一致：")
    fig(doc, J("exp1", "compare", "fig1_lda_proj_compare.jpg"), "两种实现的 LDA 投影结果对比")
    K.body(doc, "投影图中三类样本各自聚集成簇、簇间几乎不重叠，与 100% 的准确率相互印证，"
                "说明 LDA 找到的判别方向确实实现了“类内紧凑、类间疏远”的目标。两种实现给出的"
                "投影只差一个由特征向量归一化方式决定的尺度因子，判别结果不受影响。")

    # 章节结束：换页，保证后续实验另起一页
    doc.add_page_break()


def main():
    backup = "/tmp/总实验报告_Copy_备份.docx"
    shutil.copy(COPY, backup)
    print("已备份 ->", backup)

    doc = Document(COPY)
    body = doc.element.body

    # 找实验二标题作为插入点
    stop = None
    for p in doc.paragraphs:
        if p.text.strip().startswith("实验二"):
            stop = p._element
            break
    if stop is None:
        raise RuntimeError("未找到实验二标题")

    # 保留原有的封面信息表（实验一原章节里的第一张表）
    cover = None
    for el in list(body):
        if el is stop:
            break
        if el.tag == qn("w:tbl"):
            cover = el
            break

    # 在文档末尾构建新章节
    # 注意：add_paragraph 会把新元素插在 sectPr 之前，因此起点应取 sectPr 的位置
    sect = body.find(qn("w:sectPr"))
    start_len = list(body).index(sect) if sect is not None else len(list(body))
    build_section(doc)
    new_els = [el for el in list(body)[start_len:] if el.tag != qn("w:sectPr")]
    print(f"新章节元素数：{len(new_els)}，其中首元素为标题：{new_els[0].tag.split('}')[-1]}")

    # 删除旧的实验一章节（保留封面表）
    removed = 0
    for el in list(body):
        if el is stop:
            break
        if el is cover:
            continue
        body.remove(el)
        removed += 1
    print(f"已删除旧章节元素：{removed}")

    # 插入新章节
    for el in new_els:
        stop.addprevious(el)

    # 把封面表移到新标题之后
    if cover is not None:
        new_els[0].addnext(cover)

    doc.save(COPY)
    print("已写入 ->", COPY)
    print(f"插图 {NUM.fig} 张，公式 {NUM.eq} 个")


if __name__ == "__main__":
    main()
